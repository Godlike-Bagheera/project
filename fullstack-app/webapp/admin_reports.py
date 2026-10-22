"""Экспорт отчётов в Excel для админки."""
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from django.contrib.admin.views.decorators import staff_member_required
from django.db.models import Avg, Count
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from .models import Order, Product, Review, OrderItem


# ---------- Утилиты ----------

def _autosize(ws, max_width=50):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 10
        for cell in column_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, max_width)


def _style_header(ws, row=1):
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(bold=True, color="FFFFFF", size=11)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _make_response(wb, prefix):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"{prefix}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ---------- Заказы ----------

def _fill_orders_workbook(wb, orders_qs):
    orders_qs = list(orders_qs)

    ws = wb.active
    ws.title = "Заказы"
    ws.append(["№", "Пользователь", "Дата создания", "Статус", "Позиций", "Сумма, руб."])

    total_paid = Decimal("0")
    total_all = Decimal("0")
    paid_count = 0

    for o in orders_qs:
        items = list(o.items.all())
        cost = o.get_total_cost()
        total_all += cost
        if o.is_paid:
            total_paid += cost
            paid_count += 1
        ws.append([
            o.id,
            o.user.username,
            o.created_at.strftime("%d.%m.%Y %H:%M"),
            "Оплачен" if o.is_paid else "Не оплачен",
            sum(i.quantity for i in items),
            float(cost),
        ])

    _style_header(ws)
    _autosize(ws)

    ws.append([])
    ws.append(["", "", "", "Всего заказов:", len(orders_qs), ""])
    ws.append(["", "", "", "Оплачено:", paid_count, ""])
    ws.append(["", "", "", "Сумма оплаченных, руб.:", "", float(total_paid)])
    ws.append(["", "", "", "Общая сумма, руб.:", "", float(total_all)])

    for r in range(ws.max_row - 3, ws.max_row + 1):
        ws.cell(row=r, column=4).font = Font(bold=True)
        ws.cell(row=r, column=6).font = Font(bold=True)

    ws2 = wb.create_sheet("Позиции заказов")
    ws2.append(["№ заказа", "Товар", "Кол-во", "Цена, руб.", "Стоимость, руб."])
    for o in orders_qs:
        for item in o.items.all():
            ws2.append([
                o.id,
                item.product.name,
                item.quantity,
                float(item.price),
                float(item.get_cost()),
            ])
    _style_header(ws2)
    _autosize(ws2, max_width=60)
    return wb


def orders_to_excel_response(queryset):
    wb = openpyxl.Workbook()
    _fill_orders_workbook(
        wb, queryset.select_related("user").prefetch_related("items__product")
    )
    return _make_response(wb, "orders")


@staff_member_required
def export_orders_excel(request):
    qs = (
        Order.objects
        .select_related("user")
        .prefetch_related("items__product")
        .order_by("-created_at")
    )
    if (d := request.GET.get("date_from")):
        qs = qs.filter(created_at__date__gte=d)
    if (d := request.GET.get("date_to")):
        qs = qs.filter(created_at__date__lte=d)
    status = request.GET.get("status")
    if status == "paid":
        qs = qs.filter(is_paid=True)
    elif status == "unpaid":
        qs = qs.filter(is_paid=False)

    wb = openpyxl.Workbook()
    _fill_orders_workbook(wb, qs)
    return _make_response(wb, "orders")


# ---------- Детализация продаж (Кто, что, когда купил) ----------

def _fill_sales_workbook(wb, items_qs):
    ws = wb.active
    ws.title = "Детализация продаж"
    ws.append(["Дата покупки", "№ Заказа", "Покупатель", "Товар", "Кол-во", "Цена за шт., руб.", "Сумма, руб."])

    total_sum = Decimal("0")

    for item in items_qs:
        sum_cost = item.get_cost()
        total_sum += sum_cost
        ws.append([
            item.order.created_at.strftime("%d.%m.%Y %H:%M"),
            item.order.id,
            item.order.user.username,
            item.product.name,
            item.quantity,
            float(item.price),
            float(sum_cost),
        ])

    _style_header(ws)
    _autosize(ws, max_width=60)

    ws.append([])
    ws.append(["", "", "", "ИТОГО:", "", "", float(total_sum)])
    ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=7).font = Font(bold=True)

    return wb


@staff_member_required
def export_sales_excel(request):
    # Берем только оплаченные позиции
    qs = OrderItem.objects.filter(order__is_paid=True).select_related(
        "order", "order__user", "product"
    ).order_by("-order__created_at")

    if (d := request.GET.get("date_from")):
        qs = qs.filter(order__created_at__date__gte=d)
    if (d := request.GET.get("date_to")):
        qs = qs.filter(order__created_at__date__lte=d)

    wb = openpyxl.Workbook()
    _fill_sales_workbook(wb, qs)
    return _make_response(wb, "sales_details")


# ---------- Товары ----------

def _fill_products_workbook(wb, products_qs):
    ws = wb.active
    ws.title = "Товары"
    ws.append(["ID", "Название", "Категория", "Цена, руб.", "Продано, шт.", "Описание", "Ссылка"])
    for p in products_qs:
        sold = sum(i.quantity for i in p.order_items.filter(order__is_paid=True))
        ws.append([
            p.id,
            p.name,
            p.category.name if p.category else "—",
            float(p.price),
            sold,
            p.description,
            p.image_url or "",
        ])
    _style_header(ws)
    _autosize(ws, max_width=60)
    return wb


def products_to_excel_response(queryset):
    wb = openpyxl.Workbook()
    _fill_products_workbook(wb, queryset.select_related("category"))
    return _make_response(wb, "products")


@staff_member_required
def export_products_excel(request):
    wb = openpyxl.Workbook()
    _fill_products_workbook(
        wb, Product.objects.select_related("category").order_by("name")
    )
    return _make_response(wb, "products")


# ---------- Отзывы ----------

def _fill_reviews_workbook(wb, reviews_qs):
    ws = wb.active
    ws.title = "Отзывы"
    ws.append(["ID", "Товар", "Автор", "Оценка", "Дата", "Текст"])
    for r in reviews_qs:
        ws.append([
            r.id,
            r.product.name,
            r.author.username,
            r.rating,
            r.created_at.strftime("%d.%m.%Y %H:%M"),
            r.text,
        ])
    _style_header(ws)
    _autosize(ws, max_width=60)

    ws2 = wb.create_sheet("Сводка по товарам")
    ws2.append(["Товар", "Кол-во отзывов", "Средняя оценка"])
    for p in (
        Product.objects
        .annotate(cnt=Count("reviews"), avg=Avg("reviews__rating"))
        .filter(cnt__gt=0)
        .order_by("-avg")
    ):
        ws2.append([p.name, p.cnt, round(float(p.avg), 2)])
    _style_header(ws2)
    _autosize(ws2)
    return wb


def reviews_to_excel_response(queryset):
    wb = openpyxl.Workbook()
    _fill_reviews_workbook(wb, queryset.select_related("product", "author"))
    return _make_response(wb, "reviews")


@staff_member_required
def export_reviews_excel(request):
    wb = openpyxl.Workbook()
    _fill_reviews_workbook(
        wb,
        Review.objects.select_related("product", "author").order_by("-created_at"),
    )
    return _make_response(wb, "reviews")


# ---------- Страница отчётов ----------

@staff_member_required
def reports_index(request):
    return render(request, "reports.html", {
        "orders_count": Order.objects.count(),
        "products_count": Product.objects.count(),
        "reviews_count": Review.objects.count(),
    })